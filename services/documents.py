"""
Entrada/salida de documentos: extracción (.docx, .pdf, .txt, .md) y exportación
a Word (.docx) con el formato de página de la institución seleccionada.
"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader

from .institutions import Institution
from .text_utils import normalize

ALLOWED_EXTENSIONS = {"docx", "pdf", "txt", "md"}


class ExtractionError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError(f"Formato no soportado: .{ext}. Use .docx, .pdf o .txt")
    try:
        if ext == "docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n\n".join(p for p in parts if p.strip())
        elif ext == "pdf":
            reader = PdfReader(io.BytesIO(data))
            pages = [(page.extract_text() or "") for page in reader.pages]
            # Une líneas cortadas por el maquetado del PDF conservando párrafos.
            text = "\n\n".join(_unwrap_pdf(p) for p in pages)
        else:
            text = data.decode("utf-8", errors="replace")
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"No se pudo leer el archivo: {exc}") from exc
    text = normalize(text)
    if not text:
        raise ExtractionError("El archivo no contiene texto extraíble (¿PDF escaneado?)")
    return text


def _unwrap_pdf(page_text: str) -> str:
    lines = [l.rstrip() for l in page_text.splitlines()]
    out, buf = [], []
    for line in lines:
        if not line.strip():
            if buf:
                out.append(" ".join(buf)); buf = []
            continue
        if buf and buf[-1].endswith("-"):
            buf[-1] = buf[-1][:-1] + line.strip()
        else:
            buf.append(line.strip())
        if line.strip().endswith((".", ":", "?", "!")) and len(line) < 60:
            out.append(" ".join(buf)); buf = []
    if buf:
        out.append(" ".join(buf))
    return "\n".join(out)


# --------------------------------------------------------------------------- #
#  Exportación DOCX
# --------------------------------------------------------------------------- #
def _set_base_style(doc: Document, inst: Institution) -> None:
    fmt = inst.page
    style = doc.styles["Normal"]
    style.font.name = fmt.font_name
    style.font.size = Pt(fmt.font_size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), fmt.font_name)
    pf = style.paragraph_format
    pf.line_spacing = fmt.line_spacing
    pf.space_after = Pt(fmt.space_after_pt)
    pf.space_before = Pt(0)
    for section in doc.sections:
        section.top_margin = Cm(fmt.margin_top_cm)
        section.bottom_margin = Cm(fmt.margin_bottom_cm)
        section.left_margin = Cm(fmt.margin_left_cm)
        section.right_margin = Cm(fmt.margin_right_cm)
    for name in ("Heading 1", "Heading 2", "Title"):
        st = doc.styles[name]
        st.font.name = fmt.font_name
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), fmt.font_name)


def _add_page_number(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


def _body_paragraph(doc: Document, text: str, inst: Institution):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(inst.page.first_line_indent_cm)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if inst.page.justify else WD_ALIGN_PARAGRAPH.LEFT
    return p


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def build_docx(text: str, inst: Institution, title: str = "", author: str = "",
               report: dict | None = None, rewrites: list[dict] | None = None) -> bytes:
    """Genera el .docx corregido y, opcionalmente, un anexo con el informe de auditoría."""
    doc = Document()
    _set_base_style(doc, inst)
    _add_page_number(doc.sections[0])
    doc.core_properties.title = title or "Documento corregido"
    doc.core_properties.author = author or "Auditor Académico"

    if title:
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run(title); r.bold = True; r.font.size = Pt(inst.page.font_size + 2)
        if author:
            a = doc.add_paragraph(author); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

    for para in [p.strip() for p in text.split("\n") if p.strip()]:
        _body_paragraph(doc, para, inst)

    if report:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Anexo – Informe de auditoría académica", level=1)
        meta = doc.add_paragraph()
        meta.add_run(f"Norma aplicada: {inst.name} · {inst.citation_style}\n").italic = True
        meta.add_run(f"Fecha: {datetime.now():%Y-%m-%d %H:%M}")

        ai = report.get("ai", {})
        pl = report.get("plagiarism", {})
        rows = [
            ("Probabilidad de texto generado por IA", f"{ai.get('ai_probability', '—')} %",
             f"≤ {inst.ai_threshold} %"),
            ("Índice de similitud (muestral)", f"{pl.get('similarity_index', '—')} %",
             f"≤ {inst.similarity_threshold} %"),
            ("Burstiness (CV longitudes)", str(ai.get("local", {}).get("metrics", {}).get("burstiness_cv", "—")),
             "≥ 0,45 típico humano"),
            ("Perplejidad estimada", str(ai.get("local", {}).get("metrics", {}).get("perplexity_estimated", "—")), "—"),
        ]
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, label in zip(table.rows[0].cells, ("Indicador", "Resultado", "Umbral")):
            cell.text = label; cell.paragraphs[0].runs[0].bold = True; _shade(cell, "D9E2F3")
        for row in rows:
            cells = table.add_row().cells
            for c, v in zip(cells, row):
                c.text = v

        sources = pl.get("sources") or []
        if sources:
            doc.add_heading("Fuentes con similitud detectada", level=2)
            for s in sources[:15]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{s.get('title', 'Sin título')} ").bold = True
                p.add_run(f"({s.get('year') or 's.f.'}) – {s.get('source')} – similitud "
                          f"{s.get('max_similarity', s.get('similarity'))} %. {s.get('url') or ''}")

        if rewrites:
            doc.add_heading("Registro de cambios de redacción", level=2)
            t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"
            for cell, label in zip(t.rows[0].cells, ("Original", "Versión corregida")):
                cell.text = label; cell.paragraphs[0].runs[0].bold = True; _shade(cell, "D9E2F3")
            for rw in rewrites[:80]:
                c = t.add_row().cells
                c[0].text = rw.get("original", ""); c[1].text = rw.get("rewritten", "")

        note = doc.add_paragraph()
        note.add_run(ai.get("disclaimer", "")).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
